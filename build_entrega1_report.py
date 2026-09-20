from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\santy\Documents\ChatGPT\webPage")
OUT = ROOT / "output" / "docx" / "PoliTechNews_Entrega_1.docx"
IMG = ROOT / "tmp" / "report" / "mockups"
FIGMA_URL = "https://www.figma.com/design/A2QASHWYMKsV5a853lZ3dQ/PoliTechNews-%E2%80%94-Entrega-1"


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_cell_fill(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="D9D9D9", size="6"):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(underline)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.extend([fldChar1, instrText, fldChar2])


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.first_line_indent = Inches(0.5) if first_line else Inches(0)
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.add_run(item)


def add_feature(doc, title, purpose, elements, interaction):
    p = add_heading(doc, title, 2)
    p.paragraph_format.space_before = Pt(6)
    add_body(doc, purpose)
    lead = doc.add_paragraph(style="Body Text")
    lead.paragraph_format.first_line_indent = Inches(0)
    r = lead.add_run("Elementos principales. ")
    r.bold = True
    lead.add_run(elements)
    flow = doc.add_paragraph(style="Body Text")
    flow.paragraph_format.first_line_indent = Inches(0)
    r = flow.add_run("Interacción esperada. ")
    r.bold = True
    flow.add_run(interaction)


def add_figure_page(doc, number, heading, title, image_name, width, note):
    doc.add_page_break()
    add_heading(doc, heading, 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"Figura {number}")
    r.bold = True
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run(title)
    r2.italic = True
    imgp = doc.add_paragraph()
    imgp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imgp.paragraph_format.space_after = Pt(6)
    run = imgp.add_run()
    inline = run.add_picture(str(IMG / image_name), width=Inches(width))
    docPr = inline._inline.docPr
    docPr.set("descr", f"Mockup de la vista {title} de PoliTechNews")
    note_p = doc.add_paragraph(style="Body Text")
    note_p.paragraph_format.first_line_indent = Inches(0)
    note_p.paragraph_format.line_spacing = 1.0
    note_run = note_p.add_run("Nota. ")
    note_run.italic = True
    note_p.add_run(note)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.5)
section.footer_distance = Inches(0.5)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
normal.paragraph_format.space_after = Pt(0)

body = styles["Body Text"]
body.font.name = "Times New Roman"
body._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
body._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
body.font.size = Pt(12)
body.font.color.rgb = RGBColor(0, 0, 0)
body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
body.paragraph_format.space_after = Pt(0)

title_style = styles["Title"]
title_style.font.name = "Times New Roman"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
title_style.font.size = Pt(16)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(18)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name in ("Heading 1", "Heading 2", "Heading 3"):
    style = styles[name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.size = Pt(12)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.0

for sname in ("List Bullet", "List Number"):
    style = styles[sname]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    style.paragraph_format.space_after = Pt(0)

header = section.header
header.is_linked_to_previous = False
add_page_number(header.paragraphs[0])

doc.core_properties.title = "Propuesta de maquetación de PoliTechNews"
doc.core_properties.subject = "Entrega 1 del módulo Front End"
doc.core_properties.author = "Santiago Calvo Patiño"
doc.core_properties.keywords = "PoliTechNews, Front End, mockups, Figma, maquetación"

# Portada
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Propuesta de maquetación de PoliTechNews")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Santiago Calvo Patiño")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Politécnico Grancolombiano")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Front End")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("John Olarte Ramos")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Medellín")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("09/12/2026")

# Tabla de contenido
doc.add_page_break()
add_heading(doc, "Tabla de contenido", 1)
toc = [
    ("Introducción", 3),
    ("Objetivos", 4),
    ("Descripción del proyecto", 5),
    ("Diseño e información", 6),
    ("Interacción del usuario", 7),
    ("Funcionalidades principales", 8),
    ("Funcionalidades complementarias", 9),
    ("Mockup de inicio", 10),
    ("Mockup de noticias", 11),
    ("Mockup de detalle", 12),
    ("Mockup de favoritos", 13),
    ("Mockup de gestión de noticias", 14),
    ("Mockup de contacto", 15),
    ("Conclusiones", 16),
    ("Referencias", 17),
]
for label, page in toc:
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.add_run(label)
    p.add_run("\t")
    p.add_run(str(page))

# Introducción
doc.add_page_break()
add_heading(doc, "Introducción", 1)
add_body(doc, "Este informe presenta la propuesta de maquetación de PoliTechNews, una plataforma web universitaria orientada a la publicación y consulta de noticias sobre tecnología. La propuesta responde a los requisitos de la primera entrega del módulo Front End y define la base visual que posteriormente se implementará con HTML, CSS, JavaScript y fundamentos de Angular.")
add_body(doc, "El documento describe el propósito de la aplicación, su público objetivo, la navegación prevista y las funciones de cada pantalla. También incorpora los seis mockups elaborados en Figma: inicio, listado de noticias, detalle de noticia, favoritos, gestión de noticias y contacto. Las cuatro primeras vistas obligatorias indicadas por la guía académica están incluidas, y las vistas de favoritos y gestión anticipan requisitos de las siguientes etapas del proyecto.")
add_body(doc, "La maquetación mantiene una estructura consistente, con encabezado, navegación, jerarquía tipográfica, tarjetas informativas, botones de acción y pie de página. Esta coherencia facilitará que el desarrollo funcional corresponda con la propuesta visual, como exige la orientación del módulo (Olarte Ramos, 2026).")

# Objetivos
doc.add_page_break()
add_heading(doc, "Objetivos", 1)
add_heading(doc, "Objetivo general", 2)
add_body(doc, "Diseñar la propuesta visual de una plataforma web de noticias tecnológicas que permita consultar publicaciones de forma organizada, visualizar su contenido detallado, administrar favoritos, gestionar noticias y establecer contacto mediante un formulario.")
add_heading(doc, "Objetivos específicos", 2)
add_bullets(doc, [
    "Definir una identidad visual coherente y aplicable al desarrollo Front End.",
    "Organizar las noticias por categorías mediante tarjetas, filtros y búsqueda.",
    "Representar la navegación entre el listado y el detalle de cada publicación.",
    "Diseñar las vistas necesarias para guardar favoritos y gestionar noticias.",
    "Incluir un formulario de contacto con los campos y estados requeridos.",
    "Documentar los elementos y la interacción esperada de cada pantalla.",
])

# Descripción
doc.add_page_break()
add_heading(doc, "Descripción del proyecto", 1)
add_heading(doc, "Concepto", 2)
add_body(doc, "PoliTechNews es un periódico digital universitario dedicado a contenidos de inteligencia artificial, desarrollo de software, computación en la nube, datos, ciberseguridad e innovación educativa. El nombre relaciona la temática tecnológica con la identidad institucional y permite proyectar una aplicación reconocible dentro del contexto académico.")
add_heading(doc, "Público objetivo", 2)
add_body(doc, "La plataforma está dirigida principalmente a estudiantes, docentes y miembros de la comunidad universitaria interesados en comprender tendencias tecnológicas y su aplicación en procesos académicos y profesionales. La redacción, la clasificación por categorías y las acciones visibles buscan reducir el esfuerzo necesario para encontrar información relevante.")
add_heading(doc, "Alcance de la primera entrega", 2)
add_body(doc, "La primera entrega cubre la propuesta visual y la descripción funcional. No incluye todavía programación, almacenamiento local, conexión con servicios externos ni envío real del formulario. Estas funciones se implementarán en las siguientes etapas, conservando la estructura presentada en los mockups.")
p = doc.add_paragraph(style="Body Text")
p.paragraph_format.first_line_indent = Inches(0)
p.add_run("Archivo editable de Figma. ").bold = True
add_hyperlink(p, "PoliTechNews Entrega 1", FIGMA_URL)

# Diseño e información
doc.add_page_break()
add_heading(doc, "Diseño e información", 1)
add_heading(doc, "Identidad visual", 2)
add_body(doc, "La interfaz utiliza azul oscuro como color principal, cian como acento y fondos blancos o grises claros. Esta combinación separa las áreas de contenido, resalta las acciones principales y mantiene un aspecto tecnológico sobrio. Los títulos emplean una jerarquía marcada y los textos secundarios conservan contraste suficiente para facilitar la lectura.")
add_heading(doc, "Componentes reutilizables", 2)
add_bullets(doc, [
    "Encabezado con identidad de PoliTechNews y menú de navegación.",
    "Botones primarios y secundarios con estilos constantes.",
    "Tarjetas de noticias con imagen, categoría, título, resumen y acción.",
    "Campos de formulario con etiquetas visibles y estados previstos.",
    "Pie de página con información general y enlaces de navegación.",
])
add_heading(doc, "Criterios de usabilidad", 2)
add_body(doc, "Las acciones mantienen nombres breves y previsibles, mientras que el menú conserva la misma ubicación en todas las vistas. El uso de contraste, etiquetas persistentes y agrupación visual se proyecta como base para aplicar criterios de accesibilidad durante el desarrollo, de acuerdo con WCAG 2.2 (World Wide Web Consortium, 2024).")

# Interacción
doc.add_page_break()
add_heading(doc, "Interacción del usuario", 1)
add_body(doc, "El recorrido principal comienza en Inicio. Desde allí, el usuario puede abrir el catálogo de noticias, filtrar contenidos por tema y seleccionar una tarjeta para acceder al detalle. En la vista detallada podrá agregar la publicación a favoritos y regresar al listado sin perder el contexto de navegación.")
add_body(doc, "El menú también permite abrir la lista de favoritos, el panel de gestión y el formulario de contacto. Favoritos reúne las publicaciones guardadas y permite retirarlas. Gestión de noticias representa el mini CRUD solicitado para crear y eliminar registros. Contacto reúne los datos necesarios para enviar una consulta y prevé validaciones antes de confirmar el envío.")
add_heading(doc, "Mapa de navegación", 2)
add_bullets(doc, [
    "Inicio conduce a Noticias, Contacto y las publicaciones destacadas.",
    "Noticias conduce al Detalle de noticia mediante la acción Ver más.",
    "Detalle de noticia conduce a Favoritos o regresa al catálogo.",
    "Favoritos permite abrir una noticia guardada o eliminarla de la lista.",
    "Gestión permite registrar, consultar y eliminar noticias.",
    "Contacto permite diligenciar, validar y enviar una consulta.",
])
add_body(doc, "La navegación se organizó dentro de un único archivo de diseño. Figma permite crear, compartir y revisar archivos de diseño con capas y marcos editables, lo que facilita conservar las vistas relacionadas en un mismo espacio de trabajo (Figma, s. f.).")

# Funcionalidades principales
doc.add_page_break()
add_heading(doc, "Funcionalidades principales", 1)
add_feature(doc, "Inicio", "Presenta la identidad de PoliTechNews y ofrece accesos directos al contenido más relevante.", "Encabezado, sección de bienvenida, noticia principal, noticias destacadas, categorías, llamada a la acción y pie de página.", "El usuario selecciona Explorar noticias, abre una tarjeta destacada o navega hacia otra sección mediante el menú.")
add_feature(doc, "Listado de noticias", "Reúne el catálogo disponible y facilita la exploración por tema o palabra clave.", "Banner de página, buscador, filtros por categoría, contador de resultados, cuadrícula de tarjetas y paginación.", "El usuario escribe una consulta, cambia un filtro, recorre las páginas y selecciona Ver más para abrir el detalle.")
add_feature(doc, "Detalle de noticia", "Muestra la información completa de una publicación seleccionada.", "Categoría, título, fecha, imagen representativa, contenido, resumen lateral, noticias relacionadas y botones de acción.", "El usuario lee la publicación, la agrega a favoritos, comparte la información o regresa al catálogo.")

# Funcionalidades complementarias
doc.add_page_break()
add_heading(doc, "Funcionalidades complementarias", 1)
add_feature(doc, "Favoritos", "Presenta la selección personal de noticias guardadas por el usuario.", "Banner, contador, opciones de orden, tarjetas horizontales, controles para retirar elementos y acceso al detalle.", "El usuario ordena la lista, abre una publicación o elimina una noticia guardada. La persistencia se implementará posteriormente con localStorage o sessionStorage.")
add_feature(doc, "Gestión de noticias", "Representa el panel administrativo básico requerido para crear y eliminar noticias.", "Indicadores de estado, formulario de creación, carga de imagen, listado de registros, búsqueda, filtros, estados y acciones.", "El usuario completa los datos, publica o guarda un borrador y elimina registros existentes con confirmación previa.")
add_feature(doc, "Contacto", "Permite enviar preguntas, sugerencias o solicitudes al equipo de PoliTechNews.", "Datos institucionales, horarios, redes, formulario con nombre, correo, asunto, mensaje, aceptación de tratamiento de datos y preguntas frecuentes.", "El usuario completa los campos; el sistema valida los obligatorios y el formato del correo antes de mostrar la confirmación de envío.")

# Mockups
add_figure_page(doc, 1, "Mockup de inicio", "Vista de inicio de PoliTechNews", "Home.png", 4.25, "Elaboración propia en Figma. La vista integra la bienvenida, contenidos destacados, categorías y llamados a la acción.")
add_figure_page(doc, 2, "Mockup de noticias", "Listado de noticias de PoliTechNews", "Noticias.png", 4.85, "Elaboración propia en Figma. La vista organiza el catálogo mediante búsqueda, filtros, tarjetas y paginación.")
add_figure_page(doc, 3, "Mockup de detalle", "Detalle de una noticia de PoliTechNews", "Detalle.png", 4.25, "Elaboración propia en Figma. La vista presenta el contenido completo, acciones y publicaciones relacionadas.")
add_figure_page(doc, 4, "Mockup de favoritos", "Noticias favoritas de PoliTechNews", "Favoritos.png", 4.80, "Elaboración propia en Figma. La vista reúne las noticias guardadas y permite abrirlas o retirarlas de la selección.")
add_figure_page(doc, 5, "Mockup de gestión de noticias", "Gestión de noticias de PoliTechNews", "Gestion.png", 4.70, "Elaboración propia en Figma. La vista combina el formulario de creación con el listado y las acciones administrativas.")
add_figure_page(doc, 6, "Mockup de contacto", "Formulario de contacto de PoliTechNews", "Contacto.png", 4.85, "Elaboración propia en Figma. La vista reúne datos institucionales, formulario de contacto y preguntas frecuentes.")

# Conclusiones
doc.add_page_break()
add_heading(doc, "Conclusiones", 1)
add_body(doc, "La propuesta de PoliTechNews cubre los requisitos de maquetación establecidos para la primera entrega. Incluye las vistas obligatorias de inicio, listado de noticias, detalle y contacto, junto con la descripción de sus elementos e interacciones. Las vistas de favoritos y gestión amplían la propuesta y preparan la implementación de las funciones solicitadas para las siguientes etapas.")
add_body(doc, "La repetición del encabezado, las tarjetas, los botones y el pie de página establece un sistema visual coherente. Esta estructura reduce diferencias entre pantallas y facilita convertir los mockups en componentes reutilizables durante el desarrollo con HTML, CSS, JavaScript y Angular.")
add_body(doc, "El archivo de Figma mantiene las seis pantallas dentro de un mismo lienzo editable. Por tanto, la propuesta puede revisarse, ajustarse y compararse con el prototipo funcional antes de la segunda entrega. La siguiente fase deberá conservar la jerarquía visual, implementar los datos desde JSON y verificar las validaciones, los favoritos y el mini CRUD.")

# Referencias
doc.add_page_break()
add_heading(doc, "Referencias", 1)
references = [
    "Figma. (s. f.). Explore design files. Figma Learn. https://help.figma.com/hc/en-us/articles/15297425105303-Explore-design-files",
    "Norman, D. A. (2013). The design of everyday things: Revised and expanded edition. Basic Books.",
    "Olarte Ramos, J. (2026). Orientaciones para las entregas del módulo Front End Agosto 2026-1 [Documento de orientación académica]. Politécnico Grancolombiano.",
    "World Wide Web Consortium. (2024). Web Content Accessibility Guidelines (WCAG) 2.2. https://www.w3.org/TR/WCAG22/",
]
for ref in references:
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(12)
    p.add_run(ref)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
